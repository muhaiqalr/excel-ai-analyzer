import math
import os
import pandas as pd
from pathlib import Path


def parse_excel(filepath: str) -> dict:
    ext = Path(filepath).suffix.lower()

    if ext == ".csv":
        return _parse_csv(filepath)
    return _parse_xlsx(filepath)


def _parse_csv(filepath: str) -> dict:
    try:
        df = pd.read_csv(filepath)
    except UnicodeDecodeError:
        df = pd.read_csv(filepath, encoding="latin-1")

    col_types = detect_column_types(df)
    col_names = df.columns.tolist()
    missing = detect_missing_values(df)
    preview = _build_preview(df)

    sheet_info = {
        "name": "Sheet1",
        "index": 0,
        "row_count": len(df),
        "column_count": len(df.columns),
        "column_names": col_names,
        "column_types": col_types,
        "missing_values": missing,
        "preview": preview,
    }

    return {
        "total_sheets": 1,
        "total_rows": len(df),
        "total_columns": len(df.columns),
        "sheets": [sheet_info],
    }


def _parse_xlsx(filepath: str) -> dict:
    xls = pd.ExcelFile(filepath)
    sheets_info = []
    total_rows = 0
    total_columns = 0

    for idx, sheet_name in enumerate(xls.sheet_names):
        df = pd.read_excel(xls, sheet_name=sheet_name)
        col_types = detect_column_types(df)
        col_names = df.columns.tolist()
        missing = detect_missing_values(df)
        preview = _build_preview(df)

        sheets_info.append({
            "name": sheet_name,
            "index": idx,
            "row_count": len(df),
            "column_count": len(df.columns),
            "column_names": col_names,
            "column_types": col_types,
            "missing_values": missing,
            "preview": preview,
        })
        total_rows += len(df)
        total_columns = max(total_columns, len(df.columns))

    return {
        "total_sheets": len(sheets_info),
        "total_rows": total_rows,
        "total_columns": total_columns,
        "sheets": sheets_info,
    }


def detect_column_types(df: pd.DataFrame) -> dict:
    col_types = {}
    for col in df.columns:
        series = df[col]
        non_null = series.dropna()

        if len(non_null) == 0:
            col_types[col] = "text"
            continue

        if pd.api.types.is_bool_dtype(series):
            col_types[col] = "boolean"
        elif pd.api.types.is_numeric_dtype(series):
            col_types[col] = "numeric"
        elif pd.api.types.is_datetime64_any_dtype(series):
            col_types[col] = "date"
        else:
            unique_ratio = non_null.nunique() / len(non_null) if len(non_null) > 0 else 1
            if unique_ratio < 0.5 or non_null.nunique() <= 20:
                col_types[col] = "categorical"
            else:
                try:
                    pd.to_datetime(non_null.head(50), format="mixed")
                    col_types[col] = "date"
                except (ValueError, TypeError):
                    col_types[col] = "text"

    return col_types


def detect_missing_values(df: pd.DataFrame) -> dict:
    missing = {}
    for col in df.columns:
        count = int(df[col].isna().sum())
        total = len(df)
        missing[col] = {
            "count": count,
            "percentage": round(count / total * 100, 2) if total > 0 else 0.0,
        }
    return missing


def _build_preview(df: pd.DataFrame, max_rows: int = 5) -> list[dict]:
    preview_df = df.head(max_rows)
    rows = []
    for _, row in preview_df.iterrows():
        row_data = {}
        for col in df.columns:
            val = row[col]
            if pd.isna(val):
                row_data[col] = None
            elif isinstance(val, float) and val == int(val):
                row_data[col] = int(val)
            elif isinstance(val, (int, float, bool)):
                row_data[col] = val
            else:
                row_data[col] = str(val)
        rows.append(row_data)
    return rows


def analyze_sheet(filepath: str, sheet_name: str) -> dict:
    ext = Path(filepath).suffix.lower()

    if ext == ".csv":
        df = pd.read_csv(filepath)
        actual_sheet_name = "Sheet1"
    else:
        xls = pd.ExcelFile(filepath)
        if sheet_name not in xls.sheet_names:
            raise ValueError(f"Sheet '{sheet_name}' not found. Available: {xls.sheet_names}")
        df = pd.read_excel(xls, sheet_name=sheet_name)
        actual_sheet_name = sheet_name

    col_types = detect_column_types(df)
    missing = detect_missing_values(df)
    preview = _build_preview(df)

    column_details = []
    for col in df.columns:
        series = df[col]
        non_null = series.dropna()
        column_details.append({
            "name": col,
            "detected_type": col_types.get(col, "text"),
            "total_count": len(series),
            "non_null_count": int(len(non_null)),
            "missing_count": int(series.isna().sum()),
            "missing_percentage": round(int(series.isna().sum()) / len(series) * 100, 2) if len(series) > 0 else 0.0,
            "unique_count": int(non_null.nunique()),
        })

    return {
        "filename": Path(filepath).name,
        "sheet_name": actual_sheet_name,
        "row_count": len(df),
        "column_count": len(df.columns),
        "column_names": df.columns.tolist(),
        "detected_types": col_types,
        "missing_values": missing,
        "column_details": column_details,
        "preview": preview,
    }


def read_excel_data(
    filepath: str, sheet_name: str, page: int, page_size: int
) -> dict:
    ext = Path(filepath).suffix.lower()

    if ext == ".csv":
        df = pd.read_csv(filepath)
    else:
        df = pd.read_excel(filepath, sheet_name=sheet_name)

    df = df.fillna("")

    total_rows = len(df)
    total_pages = math.ceil(total_rows / page_size) if total_rows > 0 else 1

    start_idx = (page - 1) * page_size
    end_idx = min(start_idx + page_size, total_rows)
    page_df = df.iloc[start_idx:end_idx]

    columns = df.columns.tolist()
    rows = page_df.values.tolist()

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            if pd.isna(val):
                rows[i][j] = ""
            elif isinstance(val, float) and val == int(val):
                rows[i][j] = int(val)
            else:
                rows[i][j] = str(val) if not isinstance(val, (int, float, bool)) else val

    return {
        "sheet_name": sheet_name,
        "columns": columns,
        "rows": rows,
        "total_rows": total_rows,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages,
    }


def save_excel_data(filepath: str, changes: list, sheet_name: str = None, structural: list = None) -> None:
    ext = Path(filepath).suffix.lower()

    if ext == ".csv":
        _save_csv_data(filepath, changes, structural)
        return

    _save_xlsx_data(filepath, changes, sheet_name, structural)


def _save_csv_data(filepath: str, changes: list, structural: list = None) -> None:
    try:
        df = pd.read_csv(filepath)
    except UnicodeDecodeError:
        df = pd.read_csv(filepath, encoding="latin-1")

    if structural:
        for op in structural:
            if op.operation == "add_row":
                new_row = {col: "" for col in df.columns}
                if op.row_data:
                    for i, val in enumerate(op.row_data):
                        if i < len(df.columns):
                            new_row[df.columns[i]] = val
                idx = op.index if op.index is not None else len(df)
                df = pd.concat([
                    df.iloc[:idx],
                    pd.DataFrame([new_row]),
                    df.iloc[idx:]
                ]).reset_index(drop=True)
            elif op.operation == "delete_row":
                if op.index is not None and 0 <= op.index < len(df):
                    df = df.drop(df.index[op.index]).reset_index(drop=True)
            elif op.operation == "add_column":
                col_name = op.column_name or f"Column{len(df.columns) + 1}"
                idx = op.index if op.index is not None else len(df.columns)
                df.insert(idx, col_name, "")
            elif op.operation == "delete_column":
                if op.column_name and op.column_name in df.columns:
                    df = df.drop(columns=[op.column_name])
            elif op.operation == "rename_column":
                if op.column_name and op.new_name and op.column_name in df.columns:
                    df = df.rename(columns={op.column_name: op.new_name})

    for change in changes:
        col_name = change.column
        row_idx = change.row

        if col_name not in df.columns:
            raise ValueError(f"Column '{col_name}' not found in CSV file")
        if row_idx < 0 or row_idx >= len(df):
            raise ValueError(f"Row index {row_idx} out of range")

        try:
            if isinstance(change.value, str):
                try:
                    change.value = float(change.value)
                    if change.value == int(change.value):
                        change.value = int(change.value)
                except ValueError:
                    pass
        except (ValueError, TypeError):
            pass

        df.at[row_idx, col_name] = change.value

    df.to_csv(filepath, index=False)


def _save_xlsx_data(filepath: str, changes: list, sheet_name: str = None, structural: list = None) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(filepath)
    target_sheet_name = sheet_name or (changes[0].sheet_name if changes else (structural[0].sheet_name if structural else wb.sheetnames[0]))

    ws = wb[target_sheet_name]
    if ws is None:
        raise ValueError(f"Sheet '{target_sheet_name}' not found in file")

    if structural:
        for op in structural:
            if op.operation == "add_row":
                idx = op.index if op.index is not None else ws.max_row
                ws.insert_rows(idx + 1)
            elif op.operation == "delete_row":
                if op.index is not None and 0 <= op.index < ws.max_row:
                    ws.delete_rows(op.index + 1)
            elif op.operation == "add_column":
                col_name = op.column_name or f"Column{ws.max_column + 1}"
                idx = op.index if op.index is not None else ws.max_column
                ws.insert_cols(idx + 1)
                ws.cell(row=1, column=idx + 1, value=col_name)
            elif op.operation == "delete_column":
                if op.column_name:
                    for col_idx in range(1, ws.max_column + 1):
                        if ws.cell(row=1, column=col_idx).value == op.column_name:
                            ws.delete_cols(col_idx)
                            break
            elif op.operation == "rename_column":
                if op.column_name and op.new_name:
                    for col_idx in range(1, ws.max_column + 1):
                        if ws.cell(row=1, column=col_idx).value == op.column_name:
                            ws.cell(row=1, column=col_idx, value=op.new_name)
                            break

    if changes:
        for change in changes:
            if change.sheet_name != target_sheet_name:
                continue
            row_idx = change.row + 2
            col_name = change.column
            col_idx = None
            for ci in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=ci).value == col_name:
                    col_idx = ci
                    break
            if col_idx is None:
                raise ValueError(f"Column '{col_name}' not found in sheet '{target_sheet_name}'")
            if row_idx < 2 or row_idx > ws.max_row:
                raise ValueError(f"Row index {change.row} out of range")

            val = change.value
            try:
                if isinstance(val, str):
                    try:
                        val = float(val)
                        if val == int(val):
                            val = int(val)
                    except ValueError:
                        pass
            except (ValueError, TypeError):
                pass

            ws.cell(row=row_idx, column=col_idx, value=val)

    wb.save(filepath)


def create_excel_from_data(data: list[dict], sheet_name: str = "Sheet1") -> bytes:
    import io

    if not data:
        df = pd.DataFrame()
    else:
        columns = list(data[0].keys())
        rows = [list(row.values()) for row in data]
        df = pd.DataFrame(rows, columns=columns)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)

    return buffer.getvalue()
